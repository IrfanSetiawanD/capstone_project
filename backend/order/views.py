from django.db import transaction, models
from django.db.models import Count, Sum, Max
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from datetime import timedelta
from decimal import Decimal
from django.db.models import Q

from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAdminUser
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import Order, OrderItem, CustomerLoyalty, LoyaltySettings
from menu.models import Menu
from .serializers import OrderSerializer, LoyaltySettingsSerializer


# ─────────────────────────────────────────────
# ORDERS — CRUD
# ─────────────────────────────────────────────

@api_view(['POST'])
@permission_classes([AllowAny])
@transaction.atomic
def create_order(request):
    data       = request.data
    items_data = data.get('items', [])

    if not items_data:
        return Response({"error": "Items kosong"}, status=400)

    # Mendukung dua format payload:
    #   1. { customer: { phone, name }, ... }   ← format Checkout.vue (web)
    #   2. { customer_phone, customer_name, ... } ← format flat (POS)
    customer_data = data.get('customer') or {}
    phone = (customer_data.get('phone') or data.get('customer_phone', '')).strip()
    name  = (customer_data.get('name')  or data.get('customer_name',  '')).strip()

    source         = data.get('source', 'pos')
    payment_method = data.get('payment_method', 'cash')
    table_number   = data.get('table_number')
    notes          = data.get('notes', '')

    raw_payment_status = data.get('payment_status', '')

    if source == 'web':
        payment_status = 'pending'
        is_deferred    = False
    elif raw_payment_status == 'pending':
        # POS: "Makan Dulu" — bayar nanti
        payment_status = 'unpaid'
        is_deferred    = True
    else:
        # POS: bayar sekarang
        payment_status = 'paid'
        is_deferred    = False

    processed_items = []
    for item_data in items_data:
        menu_obj = get_object_or_404(Menu, id=item_data.get('menu_id'))
        if not menu_obj.is_available:
            return Response(
                {"error": f"Menu '{menu_obj.name}' sedang tidak tersedia"},
                status=400,
            )
        processed_items.append({
            "menu":  menu_obj,
            "qty":   int(item_data.get('quantity', 1)),
            "price": item_data.get('price') or menu_obj.price,
            "notes": item_data.get('notes', ''),
        })

    order = Order.objects.create(
        source              = source,
        status              = 'pending',
        payment_status      = payment_status,
        payment_method      = payment_method,
        is_deferred_payment = is_deferred,
        customer_name       = name,
        customer_phone      = phone,
        table_number        = table_number,
        notes               = notes,
    )

    OrderItem.objects.bulk_create([
        OrderItem(
            order    = order,
            menu     = item["menu"],
            quantity = item["qty"],
            price    = item["price"],
            notes    = item["notes"],
        )
        for item in processed_items
    ])

    order.recalculate_totals()

    if phone:
        order.apply_loyalty_discount()

    order.refresh_from_db()
    return Response(OrderSerializer(order).data, status=201)


@api_view(["GET"])
@permission_classes([AllowAny])
def list_orders(request):
    orders = (
        Order.objects
        .prefetch_related("items__menu")
        .order_by("-created_at")
    )
    return Response(OrderSerializer(orders, many=True).data)


@api_view(["GET"])
@permission_classes([AllowAny])
def get_order(request, pk):
    order = get_object_or_404(
        Order.objects.prefetch_related("items__menu"),
        pk=pk,
    )
    return Response(OrderSerializer(order).data)


@api_view(["GET"])
@permission_classes([AllowAny])
def active_orders_per_day(request):
    target_date = request.query_params.get("target_date")
    if not target_date:
        return Response({"error": "Parameter target_date wajib diisi"}, status=400)

    orders = (
        Order.objects
        .filter(created_at__date=target_date)
        .prefetch_related("items__menu")
        .order_by("-created_at")
    )
    return Response(OrderSerializer(orders, many=True).data)


# ─────────────────────────────────────────────
# LOYALTY — PUBLIK
# ─────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([AllowAny])
def check_loyalty_status(request):
    """
    GET /api/orders/check_loyalty_status/?phone=08xxx

    Response:
      - is_loyal          → dipakai semua Vue component
      - discount_percent  → dipakai Cart.vue (POS) & Checkout.vue (web)
      - discount_percentage → alias, nilai sama
    """
    phone = request.query_params.get("phone", "").strip()
    if not phone:
        return Response({
            "is_loyal": False,
            "discount_percent": 0,
            "discount_percentage": 0,
        })

    # Cek override manual dari admin
    loyalty_override = CustomerLoyalty.objects.filter(
        phone=phone,
        special_discount_percentage__isnull=False,
    ).first()

    if loyalty_override:
        pct = float(loyalty_override.special_discount_percentage)
        return Response({
            "is_loyal":            True,
            "discount_percent":    pct,
            "discount_percentage": pct,
        })

    # Hitung berdasarkan aturan umum
    settings = LoyaltySettings.get_settings()
    cutoff   = timezone.now() - timedelta(days=settings.period_days)
    stats    = Order.objects.filter(
        customer_phone=phone,
        status='completed',
        created_at__gte=cutoff,
    ).aggregate(jumlah=Count('id'), total=Sum('total_price'))

    is_loyal = (
        (stats['jumlah'] or 0) >= settings.min_orders
        and (stats['total'] or Decimal('0')) >= settings.min_spending
    )

    pct = float(settings.discount_percentage) if is_loyal else 0
    return Response({
        "is_loyal":            is_loyal,
        "discount_percent":    pct,
        "discount_percentage": pct,
    })


@api_view(["GET"])
@permission_classes([AllowAny])
def loyal_customers(request):
    settings = LoyaltySettings.get_settings()
    cutoff   = timezone.now() - timedelta(days=settings.period_days)

    aggregated = (
        Order.objects
        .filter(status='completed', created_at__gte=cutoff)
        .exclude(customer_phone='')
        .values('customer_phone')
        .annotate(
            order_count = Count('id'),
            total_spent = Sum('total_price'),
            last_name   = Max('customer_name'),
        )
        .order_by('-total_spent')
    )

    data = []
    for row in aggregated:
        is_loyal = (
            row['order_count'] >= settings.min_orders
            and row['total_spent'] >= settings.min_spending
        )
        data.append({
            "phone":       row['customer_phone'],
            "name":        row['last_name'] or '',
            "order_count": row['order_count'],
            "total_spent": row['total_spent'],
            "is_loyal":    is_loyal,
        })

    return Response(data)


# ─────────────────────────────────────────────
# REPORTS & DASHBOARD
# ─────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([AllowAny])
def order_reports(request):
    month = request.query_params.get("month")
    year  = request.query_params.get("year")

    qs = Order.objects.all()
    if month and year:
        qs = qs.filter(created_at__month=month, created_at__year=year)
    elif year:
        qs = qs.filter(created_at__year=year)

    top_menus = (
        OrderItem.objects.filter(order__in=qs)
        .values("menu__name")
        .annotate(total_qty=Sum("quantity"))
        .order_by("-total_qty")[:5]
    )

    return Response({
        "total_orders":  qs.count(),
        "total_revenue": qs.aggregate(total=Sum("total_price"))["total"] or 0,
        "top_menus":     list(top_menus),
    })


class DashboardStatsView(APIView):
    permission_classes = [AllowAny]

    def get(self, request):
        # Hanya order yang sudah LUNAS (payment_status=paid) untuk revenue
        paid_orders = Order.objects.filter(payment_status='paid')

        paid_stats = paid_orders.aggregate(
            total_revenue = Sum("total_price"),
            total_orders  = Count("id"),
        )

        pending   = Order.objects.filter(status="pending").count()
        completed = Order.objects.filter(status="completed").count()

        # Top 5 menu: hanya dari order yang sudah paid
        # Sertakan total_qty (jumlah porsi) dan total_revenue (omzet per menu)
        top_menus = (
            OrderItem.objects
            .filter(order__payment_status='paid')
            .values("menu__name")
            .annotate(
                total_qty     = Sum("quantity"),
                total_revenue = Sum(
                    models.ExpressionWrapper(
                        models.F("price") * models.F("quantity"),
                        output_field=models.DecimalField(),
                    )
                ),
            )
            .order_by("-total_qty")[:5]
        )

        settings_obj = LoyaltySettings.get_settings()
        cutoff       = timezone.now() - timedelta(days=settings_obj.period_days)
        loyal_count  = (
            Order.objects
            .filter(status='completed', created_at__gte=cutoff)
            .exclude(customer_phone='')
            .values('customer_phone')
            .annotate(jumlah=Count('id'), total=Sum('total_price'))
            .filter(
                jumlah__gte=settings_obj.min_orders,
                total__gte=settings_obj.min_spending,
            )
            .count()
        )

        return Response({
            "total_revenue":    paid_stats["total_revenue"] or 0,
            "total_orders":     paid_stats["total_orders"] or 0,
            "pending_orders":   pending,
            "completed_orders": completed,
            "top_menus": [
                {
                    "name":          m["menu__name"],
                    "total_qty":     m["total_qty"],
                    "total_revenue": m["total_revenue"] or 0,
                }
                for m in top_menus
            ],
            "loyal_users": loyal_count,
        })


@api_view(["GET"])
@permission_classes([AllowAny])
def admin_dashboard_daily_stats(request):
    from expenses.models import Expense

    target_date = request.query_params.get("target_date")

    revenue = 0
    if target_date:
        revenue = (
            Order.objects
            .filter(
                created_at__date=target_date,
                payment_status='paid',          # ← HANYA YANG SUDAH LUNAS
            )
            .aggregate(total=Sum("total_price"))["total"] or 0
        )

    expenses_total = 0
    if target_date:
        expenses_total = (
            Expense.objects
            .filter(date=target_date)
            .aggregate(total=Sum("amount"))["total"] or 0
        )

    return Response({
        "revenue":    revenue,
        "expenses":   expenses_total,
        "net_profit": revenue - expenses_total,
    })

@api_view(["GET"])
@permission_classes([AllowAny])
def finance_monthly_summary(request):
    """
    GET /api/orders/finance/monthly/?year=2026
    Mengembalikan pendapatan & pengeluaran per bulan dalam satu tahun.
    """
    from expenses.models import Expense
    from django.db.models import Sum
    from django.db.models.functions import TruncMonth
    import calendar

    year = int(request.query_params.get("year", timezone.now().year))

    # Revenue per bulan (paid only)
    revenue_qs = (
        Order.objects
        .filter(created_at__year=year, payment_status='paid')
        .annotate(month=TruncMonth('created_at'))
        .values('month')
        .annotate(total=Sum('total_price'))
        .order_by('month')
    )
    revenue_map = {r['month'].month: r['total'] for r in revenue_qs}

    # Expense per bulan
    expense_qs = (
        Expense.objects
        .filter(date__year=year)
        .values(month=models.ExpressionWrapper(
            models.Func('date', function='MONTH'),
            output_field=models.IntegerField()
        ))
        .annotate(total=Sum('amount'))
    )
    # Alternatif yang lebih portable (Django ORM):
    from django.db.models.functions import ExtractMonth
    expense_qs = (
        Expense.objects
        .filter(date__year=year)
        .annotate(month_num=ExtractMonth('date'))
        .values('month_num')
        .annotate(total=Sum('amount'))
    )
    expense_map = {e['month_num']: e['total'] for e in expense_qs}

    result = []
    for m in range(1, 13):
        rev = revenue_map.get(m, 0)
        exp = expense_map.get(m, 0)
        result.append({
            "month":       m,
            "month_name":  calendar.month_name[m],
            "revenue":     rev,
            "expenses":    exp,
            "net_profit":  rev - exp,
        })

    return Response({"year": year, "data": result})


@api_view(["GET"])
@permission_classes([AllowAny])
def finance_daily_summary(request):
    """
    GET /api/orders/finance/daily/?year=2026&month=6
    Mengembalikan pendapatan & pengeluaran per hari dalam satu bulan.
    """
    from expenses.models import Expense
    from django.db.models.functions import TruncDate
    import calendar

    year  = int(request.query_params.get("year",  timezone.now().year))
    month = int(request.query_params.get("month", timezone.now().month))

    # Revenue per hari (paid only)
    revenue_qs = (
        Order.objects
        .filter(created_at__year=year, created_at__month=month, payment_status='paid')
        .annotate(day=TruncDate('created_at'))
        .values('day')
        .annotate(total=Sum('total_price'))
        .order_by('day')
    )
    revenue_map = {str(r['day']): r['total'] for r in revenue_qs}

    # Expense per hari
    expense_qs = (
        Expense.objects
        .filter(date__year=year, date__month=month)
        .values('date')
        .annotate(total=Sum('amount'))
    )
    expense_map = {str(e['date']): e['total'] for e in expense_qs}

    # Isi semua hari dalam bulan tersebut
    import datetime
    days_in_month = calendar.monthrange(year, month)[1]
    result = []
    for d in range(1, days_in_month + 1):
        date_str = f"{year}-{month:02d}-{d:02d}"
        rev = revenue_map.get(date_str, 0)
        exp = expense_map.get(date_str, 0)
        result.append({
            "date":       date_str,
            "day":        d,
            "revenue":    rev,
            "expenses":   exp,
            "net_profit": rev - exp,
        })

    return Response({"year": year, "month": month, "data": result})


# ─────────────────────────────────────────────
# EXPORT
# ─────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([AllowAny])
def export_excel_report(request):
    import openpyxl
    from openpyxl.utils import get_column_letter

    month = request.query_params.get("month")
    year  = request.query_params.get("year")

    qs = Order.objects.all()
    if month and year:
        qs = qs.filter(created_at__month=month, created_at__year=year)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Pesanan"
    ws.append([
        "No. Order", "Tanggal", "Source",
        "Pelanggan", "Telepon",
        "Subtotal", "Diskon", "Total",
        "Status", "Pembayaran", "Metode",
    ])

    for order in qs.order_by('-created_at'):
        ws.append([
            order.order_number,
            order.created_at.strftime("%Y-%m-%d %H:%M"),
            order.get_source_display(),
            order.customer_name,
            order.customer_phone,
            float(order.subtotal),
            float(order.discount_amount),
            float(order.total_price),
            order.status,
            order.payment_status,
            order.payment_method or '',
        ])

    for i in range(1, 12):
        ws.column_dimensions[get_column_letter(i)].width = 18

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="laporan-{month or "all"}-{year or "all"}.xlsx"'
    )
    wb.save(response)
    return response


@api_view(["GET"])
@permission_classes([AllowAny])
def export_pdf_report(request):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    month = request.query_params.get("month")
    year  = request.query_params.get("year")

    qs = Order.objects.all()
    if month and year:
        qs = qs.filter(created_at__month=month, created_at__year=year)

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="laporan-{month or "all"}-{year or "all"}.pdf"'
    )

    p = canvas.Canvas(response, pagesize=A4)
    width, height = A4
    y = height - 50

    p.setFont("Helvetica-Bold", 14)
    p.drawString(40, y, f"Laporan Pesanan — {month or 'Semua'}/{year or 'Semua'}")
    y -= 30

    p.setFont("Helvetica-Bold", 9)
    p.drawString(40, y, "No. Order          Tanggal       Telepon          Total        Status")
    y -= 16

    p.setFont("Helvetica", 9)
    for order in qs.order_by('-created_at'):
        if y < 50:
            p.showPage()
            y = height - 50
            p.setFont("Helvetica", 9)
        line = (
            f"{order.order_number:<20} "
            f"{order.created_at.strftime('%d-%m-%Y'):<14} "
            f"{order.customer_phone:<18} "
            f"Rp{order.total_price:>12,.0f}   "
            f"{order.status}"
        )
        p.drawString(40, y, line)
        y -= 16

    p.save()
    return response


# ─────────────────────────────────────────────
# LOYALTY — ADMIN
# ─────────────────────────────────────────────

class LoyaltySettingsView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        return Response(LoyaltySettingsSerializer(LoyaltySettings.get_settings()).data)

    def put(self, request):
        settings   = LoyaltySettings.get_settings()
        serializer = LoyaltySettingsSerializer(settings, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response(serializer.data)


class LoyalCustomersView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        settings = LoyaltySettings.get_settings()
        cutoff   = timezone.now() - timedelta(days=settings.period_days)

        aggregated = (
            Order.objects
            .filter(status='completed', created_at__gte=cutoff)
            .exclude(customer_phone='')
            .values('customer_phone')
            .annotate(
                order_count   = Count('id'),
                total_belanja = Sum('total_price'),
                nama_terakhir = Max('customer_name'),
            )
            .order_by('-total_belanja')
        )

        overrides = {
            cl.phone: cl.special_discount_percentage
            for cl in CustomerLoyalty.objects.exclude(special_discount_percentage__isnull=True)
        }

        customers = []
        for row in aggregated:
            is_loyal = (
                row['order_count'] >= settings.min_orders
                and row['total_belanja'] >= settings.min_spending
            )
            customers.append({
                "phone":                       row['customer_phone'],
                "name":                        row['nama_terakhir'] or '',
                "order_count":                 row['order_count'],
                "total_spent":                 row['total_belanja'],
                "status":                      'LOYAL MEMBER' if is_loyal else 'REGULAR',
                "is_loyal":                    is_loyal,
                "special_discount_percentage": overrides.get(row['customer_phone']),
            })

        return Response({
            "settings":  LoyaltySettingsSerializer(settings).data,
            "customers": customers,
        })


class GiveSpecialPriceView(APIView):
    permission_classes = [IsAdminUser]

    def post(self, request, phone):
        discount = request.data.get('discount_percentage')
        if discount is None:
            return Response({'detail': 'discount_percentage wajib diisi'}, status=400)

        loyalty, _ = CustomerLoyalty.objects.get_or_create(
            phone=phone,
            defaults={'name': request.data.get('name', '')},
        )
        loyalty.special_discount_percentage = Decimal(str(discount))
        loyalty.save(update_fields=['special_discount_percentage', 'updated_at'])

        return Response({
            'phone':                       phone,
            'special_discount_percentage': float(loyalty.special_discount_percentage),
        })

    def delete(self, request, phone):
        CustomerLoyalty.objects.filter(phone=phone).update(
            special_discount_percentage=None
        )
        return Response(status=204)


# ─────────────────────────────────────────────
# UNPAID ORDERS & HISTORY
# ─────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([AllowAny])
def unpaid_orders(request):
    search = request.query_params.get("search", "")

    # Tangkap semua order yang belum lunas:
    #   'unpaid'  -> POS "Makan Dulu" (is_deferred_payment=True)
    #   'pending' -> Web checkout cash (belum dikonfirmasi kasir)
    qs = (
        Order.objects
        .filter(payment_status__in=["unpaid", "pending"])
        .prefetch_related("items__menu")
    )

    if search:
        qs = qs.filter(
            Q(customer_name__icontains=search)
            | Q(customer_phone__icontains=search)
            | Q(order_number__icontains=search)
        )

    return Response(
        OrderSerializer(qs.order_by("-created_at"), many=True).data
    )


@api_view(["PATCH"])
@permission_classes([AllowAny])
def pay_order(request, pk):
    order = get_object_or_404(Order, pk=pk)

    # Simpan status sebelumnya untuk trigger signal loyalty
    order._previous_status = order.status

    order.payment_status = "paid"
    order.status         = "completed"
    # Pertahankan payment_method lama jika frontend tidak kirim yang baru
    order.payment_method = request.data.get("payment_method") or order.payment_method or "cash"
    order.save()

    return Response({
        "success":      True,
        "order_number": order.order_number,
    })


@api_view(["GET"])
@permission_classes([AllowAny])
def order_history(request):
    period = request.query_params.get("period", "today")
    now    = timezone.now()

    qs = Order.objects.prefetch_related("items__menu")

    if period == "today":
        qs = qs.filter(created_at__date=now.date())
    elif period == "week":
        qs = qs.filter(created_at__gte=now - timedelta(days=7))
    elif period == "month":
        qs = qs.filter(created_at__gte=now - timedelta(days=30))

    return Response(
        OrderSerializer(qs.order_by("-created_at"), many=True).data
    )
